import pytest
from unittest.mock import patch, MagicMock
import os
import sys

# Add parent directory to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from data_helper import (
    clean_string, scrape_table, clean_default_spread, get_last_update, TABLE_CONFIGS,
    _find_sheet_by_keyword, _find_header_row, _parse_int, _parse_float, scrape_input_stats
)


# --- clean_string tests ---

class TestCleanString:
    def test_collapses_whitespace(self):
        assert clean_string("  hello   world  ") == "hello world"

    def test_strips_leading_trailing(self):
        assert clean_string("  hello  ") == "hello"

    def test_empty_string(self):
        assert clean_string("") == ""

    def test_single_word(self):
        assert clean_string("hello") == "hello"

    def test_non_string_passthrough(self):
        assert clean_string(123) == 123
        assert clean_string(None) is None
        assert clean_string(3.14) == 3.14


# --- TABLE_CONFIGS validation tests ---

class TestTableConfigs:
    """Verify that all table configs have required fields."""

    def test_all_configs_have_required_fields(self):
        required_fields = ['url', 'table_index', 'expected_columns', 'skip_rows']
        for name, config in TABLE_CONFIGS.items():
            for field in required_fields:
                assert field in config, f"Config '{name}' missing field '{field}'"

    def test_all_configs_have_header_keywords(self):
        for name, config in TABLE_CONFIGS.items():
            keywords = config.get('expected_header_keywords', [])
            assert len(keywords) > 0, f"Config '{name}' has no expected_header_keywords"

    def test_all_configs_have_min_expected_rows(self):
        for name, config in TABLE_CONFIGS.items():
            assert config.get('min_expected_rows', 0) > 0, (
                f"Config '{name}' should have min_expected_rows > 0"
            )

    def test_urls_are_valid(self):
        for name, config in TABLE_CONFIGS.items():
            assert config['url'].startswith('https://'), f"Config '{name}' URL should use HTTPS"

    def test_expected_columns_positive(self):
        for name, config in TABLE_CONFIGS.items():
            assert config['expected_columns'] > 0, f"Config '{name}' expected_columns must be positive"


# --- scrape_table tests with mocked HTML ---

def _make_mock_response(html):
    """Create a mock response object with the given HTML text."""
    mock = MagicMock()
    mock.text = html
    mock.status_code = 200
    return mock


def _build_html_table(headers, rows):
    """Build a simple HTML page with one table."""
    header_row = "<tr>" + "".join(f"<td>{h}</td>" for h in headers) + "</tr>"
    data_rows = ""
    for row in rows:
        data_rows += "<tr>" + "".join(f"<td>{v}</td>" for v in row) + "</tr>"
    return f"<html><body><table>{header_row}{data_rows}</table></body></html>"


class TestScrapeTable:
    @patch('data_helper.fetch_url_with_retry')
    def test_column_count_mismatch_returns_error(self, mock_fetch):
        """If the website changes significantly (>2 columns difference), scraper should fail loudly."""
        # ebit_growth expects 5 columns; give it 2 (difference of 3, exceeds ±2 tolerance)
        html = _build_html_table(
            ["Industry", "No of firms"],
            [["Tech", "100"]] * 60
        )
        mock_fetch.return_value = _make_mock_response(html)

        data, error = scrape_table('ebit_growth')
        assert data is None
        assert "Unexpected number of columns" in error

    @patch('data_helper.fetch_url_with_retry')
    def test_header_keyword_missing_returns_error(self, mock_fetch):
        """If column names change, header validation should catch it."""
        # ebit_growth expects 'Industry' in headers, give it different headers
        html = _build_html_table(
            ["Sector", "Count", "ROC", "Reinvest", "Growth"],
            [["Tech", "100", "10", "5", "3"]] * 60
        )
        mock_fetch.return_value = _make_mock_response(html)

        data, error = scrape_table('ebit_growth')
        assert data is None
        assert "Header validation failed" in error

    @patch('data_helper.fetch_url_with_retry')
    def test_too_few_rows_returns_error(self, mock_fetch):
        """If scraped data has suspiciously few rows, should fail."""
        html = _build_html_table(
            ["Industry", "No of firms", "ROC", "Reinvestment Rate", "Growth"],
            [["Tech", "100", "10", "5", "3"]] * 3  # Only 3 data rows
        )
        mock_fetch.return_value = _make_mock_response(html)

        data, error = scrape_table('ebit_growth')
        assert data is None
        assert "expected at least" in error

    @patch('data_helper.fetch_url_with_retry')
    def test_successful_scrape(self, mock_fetch):
        """Happy path: correct columns, headers, and enough rows."""
        rows = [[f"Industry {i}", str(i * 10), f"{i}%", f"{i * 2}%", f"{i * 3}%"]
                for i in range(1, 61)]
        html = _build_html_table(
            ["Industry", "No of firms", "ROC", "Reinvestment Rate", "Growth"],
            rows
        )
        mock_fetch.return_value = _make_mock_response(html)

        data, error = scrape_table('ebit_growth')
        assert error is None
        assert data is not None
        # 60 data rows + 1 header row = 61 total, skip_rows=1 removes header = 60
        assert len(data) == 60

    @patch('data_helper.fetch_url_with_retry')
    def test_percent_signs_stripped(self, mock_fetch):
        """Percentage signs should be removed from values."""
        rows = [["Industry 1", "10", "5%", "3%", "2%"]] * 55
        html = _build_html_table(
            ["Industry", "No of firms", "ROC", "Reinvestment Rate", "Growth"],
            rows
        )
        mock_fetch.return_value = _make_mock_response(html)

        data, error = scrape_table('ebit_growth')
        assert error is None
        # Check that % signs are gone from a data row
        for row in data:
            for val in row:
                if isinstance(val, str):
                    assert '%' not in val

    @patch('data_helper.fetch_url_with_retry')
    def test_empty_dataframe_returns_error(self, mock_fetch):
        """Empty table should return error."""
        html = "<html><body><table></table></body></html>"
        mock_fetch.return_value = _make_mock_response(html)

        data, error = scrape_table('ebit_growth')
        assert data is None
        assert error is not None

    @patch('data_helper.fetch_url_with_retry')
    def test_wrong_table_index_returns_error(self, mock_fetch):
        """If expected table index doesn't exist, should error."""
        # country_risk_premium expects table_index=1, but only 1 table exists
        html = _build_html_table(
            ["A", "B", "C", "D", "E", "F", "G", "H"],
            [["x"] * 8] * 60
        )
        mock_fetch.return_value = _make_mock_response(html)

        data, error = scrape_table('country_risk_premium')
        assert data is None
        assert "Expected at least 2 tables" in error


# --- get_last_update tests ---

class TestGetLastUpdate:
    @patch('data_helper.fetch_url_with_retry')
    def test_parses_month_year(self, mock_fetch):
        html = "<html><body><p>Last updated in January 2025</p></body></html>"
        mock_fetch.return_value = _make_mock_response(html)

        result = get_last_update("http://example.com", "Last updated in")
        assert result is not None
        assert result.year == 2025
        assert result.month == 1

    @patch('data_helper.fetch_url_with_retry')
    def test_parses_full_date(self, mock_fetch):
        html = "<html><body><p>Last updated: January 15, 2025</p></body></html>"
        mock_fetch.return_value = _make_mock_response(html)

        result = get_last_update("http://example.com", "Last updated:", date_format="%B %d, %Y")
        assert result is not None
        assert result.year == 2025
        assert result.month == 1
        assert result.day == 15

    @patch('data_helper.fetch_url_with_retry')
    def test_returns_none_when_not_found(self, mock_fetch):
        html = "<html><body><p>No update info here</p></body></html>"
        mock_fetch.return_value = _make_mock_response(html)

        result = get_last_update("http://example.com", "Last updated in")
        assert result is None


# --- DatabaseHandler tests ---

class TestDatabaseHandler:
    def test_requires_password_env_var(self):
        """DatabaseHandler should fail if DB_PASSWORD is not set."""
        from database import DatabaseHandler
        env_backup = os.environ.get('DB_PASSWORD')
        try:
            if 'DB_PASSWORD' in os.environ:
                del os.environ['DB_PASSWORD']
            with pytest.raises(KeyError):
                DatabaseHandler()
        finally:
            if env_backup is not None:
                os.environ['DB_PASSWORD'] = env_backup

    def test_context_manager_interface(self):
        """DatabaseHandler should support context manager protocol."""
        from database import DatabaseHandler
        os.environ.setdefault('DB_PASSWORD', 'test')
        handler = DatabaseHandler()
        assert hasattr(handler, '__enter__')
        assert hasattr(handler, '__exit__')


# --- _find_sheet_by_keyword tests ---

class TestFindSheetByKeyword:
    def test_exact_substring_match(self):
        sheets = ["Sheet1", "Input Stat Distributions", "Data"]
        assert _find_sheet_by_keyword(sheets, "input stat distribut") == "Input Stat Distributions"

    def test_misspelled_sheet_name(self):
        sheets = ["Sheet1", "Input Stat Distributioons", "Data"]
        result = _find_sheet_by_keyword(sheets, "input stat distribut")
        assert result == "Input Stat Distributioons"

    def test_no_match_returns_none(self):
        sheets = ["Sheet1", "Data", "Summary"]
        assert _find_sheet_by_keyword(sheets, "input stat distribut") is None

    def test_case_insensitive(self):
        sheets = ["INPUT STAT DISTRIBUTIONS"]
        assert _find_sheet_by_keyword(sheets, "input stat distribut") == "INPUT STAT DISTRIBUTIONS"


# --- _parse_int / _parse_float tests ---

class TestParseHelpers:
    def test_parse_int_with_commas(self):
        assert _parse_int("1,234") == 1234

    def test_parse_int_plain(self):
        assert _parse_int(100) == 100

    def test_parse_int_nan(self):
        assert _parse_int(float('nan')) is None

    def test_parse_int_empty(self):
        assert _parse_int("") is None

    def test_parse_float_with_percent(self):
        assert _parse_float("5.5%") == 5.5

    def test_parse_float_plain(self):
        assert _parse_float(3.14) == 3.14

    def test_parse_float_nan(self):
        assert _parse_float(float('nan')) is None

    def test_parse_float_empty(self):
        assert _parse_float("") is None


# --- scrape_input_stats tests ---

class TestScrapeInputStats:
    @patch('data_helper.fetch_url_with_retry')
    def test_returns_error_on_missing_sheet(self, mock_fetch):
        """Should return error if no matching sheet found."""
        import openpyxl
        from io import BytesIO
        wb = openpyxl.Workbook()
        wb.active.title = "WrongSheet"
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        mock_resp = MagicMock()
        mock_resp.content = buf.read()
        mock_fetch.return_value = mock_resp

        data, error = scrape_input_stats()
        assert data is None
        assert "Could not find sheet" in error

    @patch('data_helper.fetch_url_with_retry')
    def test_successful_parse(self, mock_fetch):
        """Happy path: sheet exists with correct structure and enough rows."""
        import openpyxl
        from io import BytesIO
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Input Stat Distributions"
        # Row 1: category headers
        ws.append(["", "", "Revenue Growth Rate", "", "", "Pre-tax Operating Margin",
                    "", "", "Sales to Invested Capital", "", "", "Cost of Capital",
                    "", "", "Beta", "Debt to Capital Ratio", "", "", "", ""])
        # Row 2: column headers
        ws.append(["Industry Group", "count", "First Quartile", "Median", "Third Quartile",
                    "First Quartile", "Median", "Third Quartile",
                    "First Quartile", "Median", "Third Quartile",
                    "First Quartile", "Median", "Third Quartile",
                    "median(Beta)", "First Quartile", "Median", "Third Quartile", "", ""])
        # Data rows (60 for min_rows check)
        for i in range(60):
            ws.append([f"Industry {i}", 100, 1.5, 2.0, 3.0, 4.0, 5.0, 6.0,
                       1.0, 2.0, 3.0, 8.0, 9.0, 10.0, 1.1, 0.5, 1.0, 1.5, 5.0, 10.0])
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        mock_resp = MagicMock()
        mock_resp.content = buf.read()
        mock_fetch.return_value = mock_resp

        data, error = scrape_input_stats()
        assert error is None
        assert len(data) == 60
        assert len(data[0]) == 20

    @patch('data_helper.fetch_url_with_retry')
    def test_fuzzy_sheet_name_match(self, mock_fetch):
        """Should find sheet even with a misspelled name."""
        import openpyxl
        from io import BytesIO
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Input Stat Distributioons"  # Misspelled
        ws.append(["", "", "Revenue Growth Rate", "", "", "Pre-tax Operating Margin",
                    "", "", "Sales to Invested Capital", "", "", "Cost of Capital",
                    "", "", "Beta", "Debt to Capital Ratio", "", "", "", ""])
        ws.append(["Industry Group", "count", "First Quartile", "Median", "Third Quartile",
                    "First Quartile", "Median", "Third Quartile",
                    "First Quartile", "Median", "Third Quartile",
                    "First Quartile", "Median", "Third Quartile",
                    "median(Beta)", "First Quartile", "Median", "Third Quartile", "", ""])
        for i in range(60):
            ws.append([f"Industry {i}", 100, 1.5, 2.0, 3.0, 4.0, 5.0, 6.0,
                       1.0, 2.0, 3.0, 8.0, 9.0, 10.0, 1.1, 0.5, 1.0, 1.5, 5.0, 10.0])
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        mock_resp = MagicMock()
        mock_resp.content = buf.read()
        mock_fetch.return_value = mock_resp

        data, error = scrape_input_stats()
        assert error is None
        assert len(data) == 60
